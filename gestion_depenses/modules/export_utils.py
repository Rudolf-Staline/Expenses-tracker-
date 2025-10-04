import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches

def export_to_csv(filepath, df):
    """Exporte un DataFrame en fichier CSV."""
    df.to_csv(filepath, index=False, encoding='utf-8-sig')

def export_to_excel(filepath, df):
    """Exporte un DataFrame en fichier Excel."""
    df.to_excel(filepath, index=False, sheet_name='Dépenses')

def export_to_pdf(filepath, df, report_title, totals_text):
    """Exporte le rapport (DataFrame et totaux) en fichier PDF."""
    doc = SimpleDocTemplate(filepath, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    # Titre
    story.append(Paragraph(report_title, styles['h1']))
    story.append(Paragraph("<br/><br/>", styles['Normal']))

    # Données du tableau
    data = [df.columns.to_list()] + df.values.tolist()

    # Création du tableau
    table = Table(data)
    style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ])
    table.setStyle(style)
    story.append(table)
    story.append(Paragraph("<br/><br/>", styles['Normal']))

    # Texte des totaux
    story.append(Paragraph(totals_text.replace(" | ", "<br/>"), styles['h3']))

    doc.build(story)

def export_to_docx(filepath, df, report_title, totals_text):
    """Exporte le rapport (DataFrame et totaux) en fichier DOCX."""
    doc = Document()
    doc.add_heading(report_title, level=1)

    # Ajouter le tableau
    table = doc.add_table(rows=1, cols=df.shape[1])
    table.style = 'Table Grid'

    # En-têtes
    hdr_cells = table.rows[0].cells
    for i, col_name in enumerate(df.columns):
        hdr_cells[i].text = str(col_name)

    # Lignes de données
    for index, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val)

    # Ajouter les totaux
    doc.add_paragraph("\n")
    for line in totals_text.split(" | "):
        doc.add_paragraph(line, style='Heading 3')

    doc.save(filepath)